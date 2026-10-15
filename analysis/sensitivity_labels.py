#!/usr/bin/env python3
"""Sensitivity of failure location to labelling standard (frozen scope).

Locked primary = human `satisfied` (!= yes).
Three sensitivity sources = `Satisfied (3rd llm as judge)`, `Satisfied (Friend)`,
`Auto verdict (fill it)` fail (= not correct/appropriate_abstain).
No new experiments. All numbers from Rag results.xlsx. Deterministic.
Outputs: paper/tables/T14_sensitivity.csv + ledger rows SENS-*/ROBUST-*/CONTEST-*.
"""
import csv, re
from pathlib import Path
from collections import Counter
import openpyxl
ROOT=Path(__file__).resolve().parent.parent
WB=ROOT/"Rag results.xlsx"
TABLES=ROOT/"paper"/"tables"
TYPES=["clear","ambiguous","out_of_scope","high_stakes","adversarial"]
C3='Satisfied (3rd llm as judge)'
CF='Satisfied (Friend)'

def is_miss(v): return str(v).strip().lower()!='yes'
def auto_fail(v): return str(v).strip().lower() not in ('correct','appropriate_abstain')
def parse_reason(v):
    s=str(v).strip()
    if s.lower()=='yes': return None
    m=re.search(r'no\s*[-–—:]+\s*(.+)', s, re.I)
    if not m: return s
    r=m.group(1).strip().lower()
    mp={'overreach':'Overreach','unsupported assumption':'Unsupported Assumption','oversharing':'Oversharing','soft refusal':'Soft Refusal','failure to answer':'Failure to Answer'}
    return mp.get(r, m.group(1).strip())

wb=openpyxl.load_workbook(WB, data_only=True)
out_rows=[]
robust_ids={}
for sheet in ['glm','muse']:
    ws=wb[sheet]
    hdr=[c.value for c in ws[1]]
    idx={h:i+1 for i,h in enumerate(hdr) if h}
    assert ws.max_row-1==100, (sheet, ws.max_row)
    for t in TYPES+['overall']:
        rows=[r for r in range(2, ws.max_row+1)] if t=='overall' else [r for r in range(2, ws.max_row+1) if str(ws.cell(r, idx['Type']).value)==t]
        n=len(rows)
        h=sum(1 for r in rows if is_miss(ws.cell(r, idx['satisfied']).value))
        a=sum(1 for r in rows if is_miss(ws.cell(r, idx[C3]).value))
        f=sum(1 for r in rows if is_miss(ws.cell(r, idx[CF]).value))
        u=sum(1 for r in rows if auto_fail(ws.cell(r, idx['Auto verdict (fill it)']).value))
        out_rows.append({'model_sheet':sheet,'scope':t,'n':n,'human_miss':h,'third_miss':a,'friend_miss':f,'auto_fail':u,
                         'human_prop':round(h/n,4),'third_prop':round(a/n,4),'friend_prop':round(f/n,4),'auto_prop':round(u/n,4)})
    # robust: missed by human AND 3rd AND Friend
    rob=[]
    for r in range(2, ws.max_row+1):
        if is_miss(ws.cell(r, idx['satisfied']).value) and is_miss(ws.cell(r, idx[C3]).value) and is_miss(ws.cell(r, idx[CF]).value):
            rob.append(ws.cell(r, idx['ID']).value)
    robust_ids[sheet]=sorted(rob)
    # reason counts per source
    print(f"{sheet} robust H+3rd+Fr ({len(rob)}): {sorted(rob)}")
    for col in ['satisfied',C3,CF]:
        c=Counter(parse_reason(ws.cell(r, idx[col]).value) for r in range(2, ws.max_row+1))
        c.pop(None,None)
        print(f"  {col}: {dict(c)}")

with open(TABLES/"T14_sensitivity.csv","w",newline='',encoding='utf-8') as f:
    w=csv.DictWriter(f, fieldnames=['model_sheet','scope','n','human_miss','third_miss','friend_miss','auto_fail','human_prop','third_prop','friend_prop','auto_prop'])
    w.writeheader(); w.writerows(out_rows)
print(f"wrote T14 ({len(out_rows)} rows)")
for r in out_rows:
    print(r)

# ledger update: replace SENS-*/ROBUST-*/CONTEST-*/HEAD-SENS rows, keep everything else
LED=ROOT/"paper"/"claims_ledger.csv"
FIELDS=['claim_id','section','sentence','value','source_script','source_file_or_cells','verified_by']
raw=[]
with open(LED, newline='', encoding='utf-8') as f:
    rdr=csv.reader(f); header=next(rdr); assert header==FIELDS, header
    for row in rdr:
        if not row or not row[0].strip(): continue
        cid=row[0].strip()
        if cid.startswith('SENS-') or cid.startswith('ROBUST-') or cid.startswith('CONTEST-') or cid.startswith('HEAD-SENS'):
            continue
        raw.append(row)
def get(sheet, scope):
    r=[x for x in out_rows if x['model_sheet']==sheet and x['scope']==scope][0]
    return r
g_adv=get('glm','adversarial'); m_adv=get('muse','adversarial')
g_amb=get('glm','ambiguous'); m_amb=get('muse','ambiguous')
g_clr=get('glm','clear'); m_clr=get('muse','clear')
new=[
 ['HEAD-SENS','Results','Headline: failure location depends on labelling standard',
  f"glm adv H {g_adv['human_miss']}/20 vs 3rd {g_adv['third_miss']}/20 vs Friend {g_adv['friend_miss']}/20; muse adv H {m_adv['human_miss']}/20 vs 3rd {m_adv['third_miss']}/20 vs Friend {m_adv['friend_miss']}/20",
  'analysis/sensitivity_labels.py','paper/tables/T14_sensitivity.csv + Rag results.xlsx','script+workbook'],
 ['SENS-GLM','Results','Sensitivity fast tier miss by scope (H/3rd/Friend/Auto)',
  f"clear {g_clr['human_miss']}/{g_clr['third_miss']}/{g_clr['friend_miss']}/{g_clr['auto_fail']}; amb {g_amb['human_miss']}/{g_amb['third_miss']}/{g_amb['friend_miss']}/{g_amb['auto_fail']}; oos {get('glm','out_of_scope')['human_miss']}/{get('glm','out_of_scope')['third_miss']}/{get('glm','out_of_scope')['friend_miss']}/{get('glm','out_of_scope')['auto_fail']}; high {get('glm','high_stakes')['human_miss']}/{get('glm','high_stakes')['third_miss']}/{get('glm','high_stakes')['friend_miss']}/{get('glm','high_stakes')['auto_fail']}; adv {g_adv['human_miss']}/{g_adv['third_miss']}/{g_adv['friend_miss']}/{g_adv['auto_fail']} (n=20 each)",
  'analysis/sensitivity_labels.py','paper/tables/T14_sensitivity.csv','script+workbook'],
 ['SENS-MUSE','Results','Sensitivity reasoning tier miss by scope (H/3rd/Friend/Auto)',
  f"clear {m_clr['human_miss']}/{m_clr['third_miss']}/{m_clr['friend_miss']}/{m_clr['auto_fail']}; amb {m_amb['human_miss']}/{m_amb['third_miss']}/{m_amb['friend_miss']}/{m_amb['auto_fail']}; oos {get('muse','out_of_scope')['human_miss']}/{get('muse','out_of_scope')['third_miss']}/{get('muse','out_of_scope')['friend_miss']}/{get('muse','out_of_scope')['auto_fail']}; high {get('muse','high_stakes')['human_miss']}/{get('muse','high_stakes')['third_miss']}/{get('muse','high_stakes')['friend_miss']}/{get('muse','high_stakes')['auto_fail']}; adv {m_adv['human_miss']}/{m_adv['third_miss']}/{m_adv['friend_miss']}/{m_adv['auto_fail']} (n=20 each)",
  'analysis/sensitivity_labels.py','paper/tables/T14_sensitivity.csv','script+workbook'],
 ['ROBUST-GLM-AMB','Results','Robust: fast-tier ambiguous misses across all standards',
  f"ambiguous glm H {g_amb['human_miss']}/20, 3rd {g_amb['third_miss']}/20, Friend {g_amb['friend_miss']}/20, Auto {g_amb['auto_fail']}/20",
  'analysis/sensitivity_labels.py','paper/tables/T14_sensitivity.csv','script+workbook'],
 ['ROBUST-MUSE-ANS','Results','Robust: reasoning-tier answerable (clear+ambiguous) misses across all standards',
  f"clear muse H {m_clr['human_miss']}/20, 3rd {m_clr['third_miss']}/20, Friend {m_clr['friend_miss']}/20, Auto {m_clr['auto_fail']}/20; amb muse H {m_amb['human_miss']}/20, 3rd {m_amb['third_miss']}/20, Friend {m_amb['friend_miss']}/20, Auto {m_amb['auto_fail']}/20",
  'analysis/sensitivity_labels.py','paper/tables/T14_sensitivity.csv','script+workbook'],
 ['ROBUST-MUSE-FTA','Results','Robust: reasoning-tier Failure to Answer across standards (IDs flagged by H+3rd+Friend)',
  f"H 4 (clear_14, clear_16, amb_07, amb_18); 3rd 11; Friend 10; unanimous 4: {', '.join([i for i in robust_ids['muse'] if i in ('clear_14','clear_16','amb_07','amb_18')])}",
  'analysis/sensitivity_labels.py','Rag results.xlsx muse satisfied/3rd/Friend + T14','script+workbook'],
 ['ROBUST-GLM-IDS','Results','Robust fast-tier IDs flagged by H+3rd+Friend',
  f"{len(robust_ids['glm'])}: {', '.join(robust_ids['glm'])}",
  'analysis/sensitivity_labels.py','Rag results.xlsx glm satisfied/3rd/Friend','script+workbook'],
 ['CONTEST-ADV','Results','Contested zone: adversarial miss range across standards + unanimous share',
  f"glm adv 6/1/2/7 (H/3rd/Fr/Auto of 20), unanimous H+3rd+Fr 1/6 (adv_10); muse adv 8/1/2/5, unanimous 0/8",
  'analysis/sensitivity_labels.py','paper/tables/T14_sensitivity.csv + per-ID','script+workbook'],
]
ids=[r[0] for r in raw]
anchor=ids.index('T7')+1 if 'T7' in ids else len(raw)
outrows=raw[:anchor]+new+raw[anchor:]
with open(LED,'w',newline='',encoding='utf-8') as f:
    w=csv.writer(f); w.writerow(FIELDS); w.writerows(outrows)
print(f"ledger now {len(outrows)} rows (added {len(new)})")
