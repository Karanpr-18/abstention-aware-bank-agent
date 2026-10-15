#!/usr/bin/env python3
"""Judge check: human vs self-judge (Auto) vs Friend (classmate) vs 3rd LLM.
Binary Yes vs No. Metrics: percent agreement, Cohen kappa, Gwet AC1.
Outputs: paper/tables/T13_judge_check.csv + notes/label_sources_by_id.csv
All from Rag results.xlsx. No hallucination.

Schema (2026-10-11): BOTH sheets carry `Satisfied (3rd llm as judge)` and
`Satisfied (Friend)`. The old glm-only `Satisfied (YES/NO)` column was
renamed to the 3rd-LLM column with identical values (verified 100/100
binary + 0 string disagreements), so YES/NO handling below is legacy-only.
Also refreshes the JUDGE-* numeric rows in paper/claims_ledger.csv from
this script's output (numbers by script, labels by hand).
"""
import csv, re
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parent.parent
WB=ROOT/"Rag results.xlsx"

def binyes(v): return str(v).strip().lower()=='yes'
def autopass(v): return str(v).strip().lower() in ('correct','appropriate_abstain')
def norm_reason(v):
    s=str(v).strip()
    if s.lower()=='yes': return 'Yes'
    m=re.search(r'no\s*[-–—:]+\s*(.+)', s, re.I)
    if not m: return s
    r=m.group(1).strip().lower()
    mp={'overreach':'Overreach','unsupported assumption':'Unsupported Assumption','oversharing':'Oversharing','soft refusal':'Soft Refusal','failure to answer':'Failure to Answer'}
    return mp.get(r, m.group(1).strip())

def metrics(a,b):
    n=len(a)
    pa=sum(1 for x,y in zip(a,b) if x==y)/n
    p1=sum(a)/n; p2=sum(b)/n
    pe_k=p1*p2+(1-p1)*(1-p2)
    k=(pa-pe_k)/(1-pe_k) if pe_k!=1 else 0.0
    pi=(p1+p2)/2
    pe_g=2*pi*(1-pi)
    ac1=(pa-pe_g)/(1-pe_g) if pe_g!=1 else 0.0
    return pa,k,ac1

C3='Satisfied (3rd llm as judge)'
CF='Satisfied (Friend)'
CYN='Satisfied (YES/NO)'  # legacy glm-only name, retired 2026-10-11

wb=openpyxl.load_workbook(WB, data_only=True)
rows_out=[]
adj=[]
for s in ['glm','muse']:
    ws=wb[s]
    hdr=[ws.cell(1,c).value for c in range(1, ws.max_column+1)]
    idx={h:i+1 for i,h in enumerate(hdr) if h}
    n=ws.max_row-1
    assert n==100, (s, n)
    human=[binyes(ws.cell(r, idx['satisfied']).value) for r in range(2, ws.max_row+1)]
    auto=[autopass(ws.cell(r, idx['Auto verdict (fill it)']).value) for r in range(2, ws.max_row+1)]
    pairs=[('human-vs-selfjudge(Auto)', human, auto, 'secondary: original same-model first pass')]
    has3=C3 in idx
    hasF=CF in idx
    t3=fr=None
    if has3:
        t3=[binyes(ws.cell(r, idx[C3]).value) for r in range(2, ws.max_row+1)]
        pairs.append(('human-vs-3rdLLM(5reason)', human, t3, 'primary: independent 5-reason check as supplied'))
    if hasF:
        fr=[binyes(ws.cell(r, idx[CF]).value) for r in range(2, ws.max_row+1)]
        pairs.append(('human-vs-classmate(Friend)', human, fr, 'primary: classmate check as supplied; blinding protocol unlogged'))
    if has3 and hasF:
        pairs.append(('3rdLLM-vs-classmate', t3, fr, 'cross-check'))
    if has3:
        pairs.append(('selfjudge-vs-3rdLLM', auto, t3, 'secondary comparison'))
    if hasF:
        pairs.append(('selfjudge-vs-classmate', auto, fr, 'secondary comparison'))
    if CYN in idx:  # legacy only; retired schema
        yn=[binyes(ws.cell(r, idx[CYN]).value) for r in range(2, ws.max_row+1)]
        pairs.append(('human-vs-indep(YESNO-5reason)', human, yn, 'LEGACY column; retired 2026-10-11'))
    for name,a,b,note in pairs:
        pa,k,ac1=metrics(a,b)
        ag=sum(1 for x,y in zip(a,b) if x==y)
        rows_out.append({'model_sheet':s,'comparison':name,'n':n,'agree':ag,'agree_prop':round(pa,4),'kappa':round(k,4),'gwet_ac1':round(ac1,4),'note':note})
    # per-ID label-source rows (no adjudication)
    for r in range(2, ws.max_row+1):
        idv=ws.cell(r, idx['ID']).value
        typ=ws.cell(r, idx['Type']).value
        gold=ws.cell(r, idx['Expected answer (gold - researcher only)']).value
        sat=ws.cell(r, idx['satisfied']).value
        auto_v=ws.cell(r, idx['Auto verdict (fill it)']).value
        hb='Yes' if binyes(sat) else f"No-{norm_reason(sat)}"
        ab='Pass' if autopass(auto_v) else f"Fail({auto_v})"
        e1=ws.cell(r, idx[C3]).value if has3 else ''
        e2=ws.cell(r, idx[CF]).value if hasF else ''
        e1b='Yes' if binyes(e1) else (f"No-{norm_reason(e1)}" if has3 else '')
        e2b='Yes' if binyes(e2) else (f"No-{norm_reason(e2)}" if hasF else '')
        adj.append({'sheet':s,'ID':idv,'Type':typ,'gold':gold,'human':hb,'selfjudge':ab,'thirdLLM':e1b,'classmate':e2b})

with open(ROOT/'paper'/'tables'/'T13_judge_check.csv','w',newline='',encoding='utf-8') as f:
    w=csv.DictWriter(f, fieldnames=['model_sheet','comparison','n','agree','agree_prop','kappa','gwet_ac1','note'])
    w.writeheader(); w.writerows(rows_out)
with open(ROOT/'notes'/'label_sources_by_id.csv','w',newline='',encoding='utf-8') as f:
    w=csv.DictWriter(f, fieldnames=['sheet','ID','Type','gold','human','selfjudge','thirdLLM','classmate'])
    w.writeheader(); w.writerows(adj)
print(f"wrote T13 ({len(rows_out)} rows) + label_sources ({len(adj)} rows)")
for r in rows_out:
    print(r)

# ---- refresh numeric JUDGE rows in claims ledger from this script's output ----
# Robust raw-row handling: filter on column 0 text; never drop other rows.
LED=ROOT/'paper'/'claims_ledger.csv'
FIELDS=['claim_id','section','sentence','value','source_script','source_file_or_cells','verified_by']
raw=[]
with open(LED, newline='', encoding='utf-8') as f:
    rdr=csv.reader(f)
    header=next(rdr)
    assert header==FIELDS, header
    for row in rdr:
        if not row or not row[0].strip():
            continue
        cid=row[0].strip()
        if cid.startswith('JUDGE-P-') or cid.startswith('JUDGE-X-'):
            continue
        raw.append(row)
def get(sheet, comp):
    r=[x for x in rows_out if x['model_sheet']==sheet and x['comparison']==comp][0]
    return f"{r['agree']}/{r['n']}={r['agree_prop']*100:.1f}% kappa {r['kappa']:.2f} AC1 {r['gwet_ac1']:.2f}"
def T(sheet, comp): return [x for x in rows_out if x['model_sheet']==sheet and x['comparison']==comp][0]
g3=T('glm','human-vs-3rdLLM(5reason)'); gF=T('glm','human-vs-classmate(Friend)'); gx=T('glm','3rdLLM-vs-classmate')
m3=T('muse','human-vs-3rdLLM(5reason)'); mF=T('muse','human-vs-classmate(Friend)'); mx=T('muse','3rdLLM-vs-classmate')
new=[
 {'claim_id':'JUDGE-P-GLM3','section':'Results','sentence':'Primary independent 5-reason check fast tier human-vs-3rdLLM','value':get('glm','human-vs-3rdLLM(5reason)'),'source_script':'analysis/judge_check.py','source_file_or_cells':'paper/tables/T13_judge_check.csv + Rag results.xlsx glm Satisfied (3rd llm as judge)','verified_by':'script+workbook'},
 {'claim_id':'JUDGE-P-GLMF','section':'Results','sentence':'Primary classmate check fast tier human-vs-Friend','value':get('glm','human-vs-classmate(Friend)'),'source_script':'analysis/judge_check.py','source_file_or_cells':'paper/tables/T13_judge_check.csv + Rag results.xlsx glm Satisfied (Friend)','verified_by':'script+workbook'},
 {'claim_id':'JUDGE-P-MUSE3','section':'Results','sentence':'Primary independent 5-reason check reasoning human-vs-3rdLLM','value':get('muse','human-vs-3rdLLM(5reason)'),'source_script':'analysis/judge_check.py','source_file_or_cells':'paper/tables/T13_judge_check.csv + Rag results.xlsx muse Satisfied (3rd llm as judge)','verified_by':'script+workbook'},
 {'claim_id':'JUDGE-P-MUSEF','section':'Results','sentence':'Primary classmate check reasoning human-vs-Friend','value':get('muse','human-vs-classmate(Friend)'),'source_script':'analysis/judge_check.py','source_file_or_cells':'paper/tables/T13_judge_check.csv + Rag results.xlsx muse Satisfied (Friend)','verified_by':'script+workbook'},
 {'claim_id':'JUDGE-X-GLM','section':'Results','sentence':'Cross-check 3rdLLM-vs-classmate fast tier','value':get('glm','3rdLLM-vs-classmate'),'source_script':'analysis/judge_check.py','source_file_or_cells':'paper/tables/T13_judge_check.csv','verified_by':'script+workbook'},
 {'claim_id':'JUDGE-X-MUSE','section':'Results','sentence':'Cross-check 3rdLLM-vs-classmate reasoning','value':get('muse','3rdLLM-vs-classmate'),'source_script':'analysis/judge_check.py','source_file_or_cells':'paper/tables/T13_judge_check.csv','verified_by':'script+workbook'},
]
# keep ledger order: insert new rows where old JUDGE-P rows sat (after JUDGE-SELF)
ids=[r[0] for r in raw]
anchor=ids.index('JUDGE-SELF')+1 if 'JUDGE-SELF' in ids else len(raw)
outrows=raw[:anchor]+[[d['claim_id'],d['section'],d['sentence'],d['value'],d['source_script'],d['source_file_or_cells'],d['verified_by']] for d in new]+raw[anchor:]
with open(LED,'w',newline='',encoding='utf-8') as f:
    w=csv.writer(f, quoting=csv.QUOTE_MINIMAL)
    w.writerow(FIELDS)
    w.writerows(outrows)
print(f"ledger refreshed: {len(outrows)} rows")
