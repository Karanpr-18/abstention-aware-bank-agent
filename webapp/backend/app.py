"""
FastAPI backend for the Financial RAG-Agent webapp.
- Serves the React build (../frontend/dist) in production.
- /api/chat      : public chat — NEVER receives gold labels.
- /api/questions : researcher-only gold set (frontend eval page only).
- /api/eval/run  : takes qid, looks up question text SERVER-SIDE and sends
                   ONLY the text to the CrewAI agent. Gold never touches the agent.
Run: .venv/bin/uvicorn webapp.backend.app:app --port 8000
"""
import os, sys, json, time, threading, uuid
from pathlib import Path
from fastapi import FastAPI, HTTPException
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from dotenv import load_dotenv

load_dotenv()
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from src.crew_agent import run_crew_query, MODELS, BY_MODEL, DEFAULT_MODEL

app = FastAPI(title="Banking Support Agent — Reliability Testbed")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

GOLD = json.loads((ROOT / "data" / "queries_gold.json").read_text())
EXTRA = json.loads((ROOT / "data" / "queries_extra.json").read_text()) if (ROOT / "data" / "queries_extra.json").exists() else []
ALL_Q = GOLD + EXTRA
BY_ID = {q["id"]: q for q in ALL_Q}

def batches():
    """5 batches x 20: round-robin across types so every batch holds 4 of each type."""
    types = ["clear", "ambiguous", "out_of_scope", "high_stakes", "adversarial"]
    per_type = {t: [q for q in ALL_Q if q["type"] == t] for t in types}
    n = max(len(v) for v in per_type.values()) // 4
    out = []
    for i in range(n):
        b = [q for t in types for q in per_type[t][i * 4:i * 4 + 4]]
        out.append(b)
    return out

class ChatIn(BaseModel):
    message: str
    mode: str = "A"
    threshold: float = 0.5
    model: str = DEFAULT_MODEL

class EvalIn(BaseModel):
    qid: str
    mode: str = "A"
    threshold: float = 0.5
    model: str = DEFAULT_MODEL

def _check_mode(mode: str):
    if mode not in ("A", "B", "C") and not mode.startswith("B@"):
        raise HTTPException(400, "mode must be A, B or C")

def _check_model(model: str):
    if model not in BY_MODEL:
        raise HTTPException(400, f"unknown model '{model}' — pick from /api/models")

@app.get("/api/health")
def health():
    return {"ok": True, "questions": len(ALL_Q), "batches": len(batches()),
            "groq_key": bool(os.environ.get("GROQ_API_KEY")),
            "gemini_key": bool(os.environ.get("GEMINI_API_KEY")),
            "buyer_identity": os.environ.get("ANTSEED_BUYER_IDENTITY", "")}

@app.get("/api/models")
def models():
    return {"default": DEFAULT_MODEL, "models": MODELS}

@app.get("/api/questions")
def questions(batch: int | None = None):
    if batch is not None:
        bl = batches()
        if not 1 <= batch <= len(bl):
            raise HTTPException(404, "unknown batch")
        return bl[batch - 1]
    return ALL_Q

@app.get("/api/batches")
def batch_list():
    return [{"n": i + 1, "size": len(b),
             "types": {t: sum(1 for q in b if q["type"] == t) for t in ("clear", "ambiguous", "out_of_scope", "high_stakes", "adversarial")}}
            for i, b in enumerate(batches())]

class BatchIn(BaseModel):
    batch: int = 1
    mode: str = "A"
    threshold: float = 0.5
    model: str = DEFAULT_MODEL
    delay: float = 2.0
    workers: int = 4

class JudgeIn(BaseModel):
    batch: int = 1
    mode: str = "A"
    model: str = DEFAULT_MODEL
    workers: int = 4

_BATCH_DONE: dict = {}

@app.post("/api/eval/batch")
def eval_batch(body: BatchIn):
    """Run one 20-question batch as a background job. Progress polls via /api/jobs.
    Only question TEXT reaches the agent; gold stays server-side."""
    _check_mode(body.mode)
    _check_model(body.model)
    bl = batches()
    if not 1 <= body.batch <= len(bl):
        raise HTTPException(404, "unknown batch")
    qs = bl[body.batch - 1]

    jid = uuid.uuid4().hex[:12]
    with _JOB_LOCK:
        _JOBS[jid] = {"status": "running", "kind": "batch", "batch": body.batch,
                      "events": [{"stage": "presented", "detail": f"Batch {body.batch} convened — {len(qs)} matters.", "at": time.time()}],
                      "result": None, "error": None, "done_rows": []}

    def _emit(stage, detail=""):
        with _JOB_LOCK:
            job = _JOBS.get(jid)
            if not job:
                return
            if stage == "step" and job["events"]:
                job["events"][-1]["detail"] = detail or job["events"][-1]["detail"]
            else:
                job["events"].append({"stage": stage, "detail": detail, "at": time.time()})

    def _one(idx_q):
        idx, q = idx_q
        _emit("retrieving", f"Matter {idx + 1}/{len(qs)} · {q['id']} — consulting the record…")

        def _tagged(stage, detail=""):
            if stage == "step" or not detail or detail.startswith(f"[{q['id']}]"):
                _emit(stage, detail)
            else:
                _emit(stage, f"[{q['id']}] {detail}")

        try:
            r = run_crew_query(q["query"], mode=body.mode, threshold=body.threshold,
                               model=body.model, progress=_tagged)
            row = (idx, {"qid": q["id"], "qtype": q["type"], "query": q["query"],
                         "gold": q["gold"], "note": q.get("note", ""),
                         "answer": r["answer"], "abstained": r["abstained"],
                         "confidence": round(r["confidence"], 3), "mode": r["mode"],
                         "latency": round(r["latency"], 2), "model": r["model"],
                         "buyer": r.get("buyer"),
                         "sources": [{"file": m.get("source_file"), "sim": round(s, 3)} for _, s, m in r["chunks"][:4]]})
        except RuntimeError as e:
            if str(e).startswith("RATE_LIMITED"):
                raise
            row = (idx, {"qid": q["id"], "qtype": q["type"], "query": q["query"],
                         "gold": q["gold"], "note": q.get("note", ""), "error": str(e)[:300]})
        with _JOB_LOCK:
            job = _JOBS.get(jid)
            if job:
                job.setdefault("rows_by_idx", {})[idx] = row[1]
                job["done_rows"] = [job["rows_by_idx"][k] for k in sorted(job["rows_by_idx"])]
        time.sleep(body.delay)
        return row

    def _work():
        rows: dict = {}
        try:
            from concurrent.futures import ThreadPoolExecutor, as_completed
            workers = max(1, min(int(body.workers), len(qs)))
            _emit("retrieving", f"{workers} clerks deployed on {len(qs)} matters…")
            with ThreadPoolExecutor(max_workers=workers) as pool:
                futs = {pool.submit(_one, (i, q)): i for i, q in enumerate(qs)}
                try:
                    for fut in as_completed(futs):
                        idx, row = fut.result()  # re-raises RATE_LIMITED to halt the batch
                        rows[idx] = row
                except Exception:
                    for f in futs:
                        f.cancel()
                    raise
            ordered = [rows[k] for k in sorted(rows)]
            payload = {"batch": body.batch, "mode": body.mode, "model": body.model, "rows": ordered}
            with _JOB_LOCK:
                _JOBS[jid]["result"] = payload
                _JOBS[jid]["status"] = "done"
                _BATCH_DONE[(body.batch, body.mode, body.model)] = payload
        except RuntimeError as e:
            with _JOB_LOCK:
                _JOBS[jid]["error"] = {"code": 429 if str(e).startswith("RATE_LIMITED") else 500,
                                       "message": str(e).replace("RATE_LIMITED: ", "")}
                _JOBS[jid]["status"] = "error"
                _JOBS[jid]["done_rows"] = [rows[k] for k in sorted(rows)]
        except Exception as e:
            with _JOB_LOCK:
                _JOBS[jid]["error"] = {"code": 500, "message": str(e)[:300]}
                _JOBS[jid]["status"] = "error"
                _JOBS[jid]["done_rows"] = [rows[k] for k in sorted(rows)]

    threading.Thread(target=_work, daemon=True).start()
    return {"job_id": jid, "batch": body.batch, "size": len(qs)}

@app.post("/api/eval/batch/judge")
def eval_batch_judge(body: JudgeIn):
    """Mimo satisfaction pass over STORED batch rows. Writes satisfy_llm per row.
    Progress polls via /api/jobs. Runs in parallel workers; needs the bridge up."""
    _check_mode(body.mode)
    _check_model(body.model)
    payload = _BATCH_DONE.get((body.batch, body.mode, body.model))
    if not payload:
        raise HTTPException(404, "batch not run yet with this mode/model — convene it first")
    from check_satisfaction import judge as mimo_judge

    jid = uuid.uuid4().hex[:12]
    with _JOB_LOCK:
        _JOBS[jid] = {"status": "running", "kind": "judge", "batch": body.batch,
                      "events": [{"stage": "presented",
                                  "detail": f"Judging batch {body.batch} — {len(payload['rows'])} matters.",
                                  "at": time.time()}],
                      "result": None, "error": None, "done_rows": []}

    def _one(idx_row):
        idx, row = idx_row
        if row.get("error") or not row.get("answer"):
            verdict = "NOT SATISFIED — Unheard (no answer captured)"
        else:
            try:
                ok, reason = mimo_judge(row.get("gold", ""), row.get("answer", ""),
                                        question=row.get("query", ""))
                verdict = "SATISFIED" if ok else f"NOT SATISFIED — {reason}"
            except Exception as e:
                verdict = f"JUDGE FAILED — {str(e)[:120]}"
        with _JOB_LOCK:
            job = _JOBS.get(jid)
            if job:
                job.setdefault("rows_by_idx", {})[idx] = verdict
                job["done_rows"] = [{"qid": payload["rows"][k].get("qid"), "verdict": v}
                                    for k, v in sorted(job["rows_by_idx"].items())]
        return idx, verdict

    def _work():
        try:
            from concurrent.futures import ThreadPoolExecutor, as_completed
            workers = max(1, min(int(body.workers), len(payload["rows"])))
            with ThreadPoolExecutor(max_workers=workers) as pool:
                futs = [pool.submit(_one, (i, r)) for i, r in enumerate(payload["rows"])]
                verdicts = dict(f.result() for f in as_completed(futs))
            for i, v in verdicts.items():
                payload["rows"][i]["satisfy_llm"] = v
            sat = sum(1 for v in verdicts.values() if v == "SATISFIED")
            with _JOB_LOCK:
                _JOBS[jid]["result"] = {"batch": body.batch, "satisfied": sat,
                                        "total": len(verdicts)}
                _JOBS[jid]["status"] = "done"
        except Exception as e:
            with _JOB_LOCK:
                _JOBS[jid]["error"] = {"code": 500, "message": str(e)[:300]}
                _JOBS[jid]["status"] = "error"

    threading.Thread(target=_work, daemon=True).start()
    return {"job_id": jid, "batch": body.batch}

@app.get("/api/eval/batch/{n}/xlsx")
def batch_xlsx(n: int, mode: str = "A", model: str = DEFAULT_MODEL):
    """Download one completed batch as Excel: question, agent answer, expected answer."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from fastapi.responses import Response
    payload = _BATCH_DONE.get((n, mode, model))
    if not payload:
        raise HTTPException(404, "batch not run yet with this mode/model — convene it first")
    wb = Workbook()
    ws = wb.active
    ws.title = f"Batch {n}"
    head = ["No", "ID", "Type", "Question", "Expected answer (gold — researcher only)",
            "Agent answer", "Abstained", "Auto verdict", "Satisfy (LLM)", "Confidence", "Latency (s)", "Model", "Buyer", "Sources"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0C2B23")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, r in enumerate(payload["rows"], 1):
        ws.append([i, r.get("qid"), r.get("qtype"), r.get("query"), r.get("gold"),
                   r.get("answer", r.get("error", "")), r.get("abstained", ""),
                   _auto_verdict(r), r.get("satisfy_llm", ""), r.get("confidence", ""), r.get("latency", ""),
                   str(r.get("model", "")).split("/")[-1].split("@")[-1], r.get("buyer", ""),
                   "; ".join(s.get("file", "") for s in r.get("sources", []))])
    for col in ("D", "E", "F"):
        ws.column_dimensions[col].width = 48
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=batch{n}_{mode}_{str(model).split('/')[-1].split('@')[-1]}.xlsx"})

def _auto_verdict(r: dict) -> str:
    if r.get("error"):
        return "unheard"
    gold, ans, abst = r.get("gold", ""), r.get("answer", ""), r.get("abstained", False)
    a = ans.lower()[:300]
    decline = ["don't know", "do not know", "not sure", "uncertain", "unable to", "not able to",
               "cannot answer", "can't answer", "cannot determine", "can't determine", "cannot confirm",
               "outside my scope", "beyond my scope", "human specialist", "human advisor",
               "specialist", "escalat", "refer you", "decline to", "insufficient information"]
    denied = abst or any(k in a for k in decline)
    if gold in ("abstain", "refuse", "escalate"):
        return "appropriate_abstain" if denied else "false_confidence"
    if gold == "clarify":
        if any(k in a for k in ["clarif", "missing", "need", "which", "provide"]) or denied:
            return "appropriate_abstain" if denied else "correct"
        return "false_confidence"
    return "inappropriate_abstain" if denied else "correct"

@app.post("/api/chat")
def chat(body: ChatIn):
    _check_mode(body.mode)
    _check_model(body.model)
    if not body.message.strip():
        raise HTTPException(400, "empty message")
    # gold is NEVER accepted here — only raw user text reaches the agent.
    # Long crew runs execute as a background job; the UI polls its timeline.
    return _launch(body.message.strip(), body.mode, body.threshold, body.model, extra={})

@app.post("/api/eval/run")
def eval_run(body: EvalIn):
    """Researcher harness: qid -> question TEXT ONLY to agent. Gold stays server-side for display/scoring."""
    _check_mode(body.mode)
    _check_model(body.model)
    q = BY_ID.get(body.qid)
    if not q:
        raise HTTPException(404, "unknown qid")
    return _launch(q["query"], body.mode, body.threshold, body.model,
                   extra={"qid": q["id"], "qtype": q["type"], "gold": q["gold"], "note": q.get("note", "")})

_JOBS: dict = {}
_JOB_LOCK = threading.Lock()

def _launch(question: str, mode: str, threshold: float, model: str, extra: dict):
    jid = uuid.uuid4().hex[:12]
    with _JOB_LOCK:
        _JOBS[jid] = {"status": "running", "events": [{"stage": "presented", "detail": "Your enquiry is presented at the counter.",
                                                        "at": time.time()}], "result": None, "error": None}
        if len(_JOBS) > 100:
            _JOBS.pop(next(iter(_JOBS)))
    def _emit(stage, detail=""):
        if stage == "step":
            with _JOB_LOCK:
                job = _JOBS.get(jid)
                if job and job["events"]:
                    job["events"][-1]["detail"] = detail or job["events"][-1]["detail"]
            return
        with _JOB_LOCK:
            job = _JOBS.get(jid)
            if job:
                job["events"].append({"stage": stage, "detail": detail, "at": time.time()})
    def _work():
        try:
            r = run_crew_query(question, mode=mode, threshold=threshold, model=model, progress=_emit)
            payload = {"query": question,
                       "answer": r["answer"], "abstained": r["abstained"],
                       "confidence": round(r["confidence"], 3), "mode": r["mode"],
                       "latency": round(r["latency"], 2), "model": r["model"],
                       "buyer": r.get("buyer"),
                       "sources": [{"file": m.get("source_file"), "sim": round(s, 3)} for _, s, m in r["chunks"][:4]],
                       **extra}
            with _JOB_LOCK:
                _JOBS[jid]["result"] = payload
                _JOBS[jid]["status"] = "done"
        except RuntimeError as e:
            code = 429 if str(e).startswith("RATE_LIMITED") else 500
            msg = str(e).replace("RATE_LIMITED: ", "")
            with _JOB_LOCK:
                _JOBS[jid]["error"] = {"code": code, "message": msg}
                _JOBS[jid]["status"] = "error"
        except Exception as e:
            with _JOB_LOCK:
                _JOBS[jid]["error"] = {"code": 500, "message": str(e)[:300]}
                _JOBS[jid]["status"] = "error"
    threading.Thread(target=_work, daemon=True).start()
    return {"job_id": jid}

@app.get("/api/jobs/{jid}")
def job_state(jid: str):
    with _JOB_LOCK:
        job = _JOBS.get(jid)
        if not job:
            raise HTTPException(404, "unknown job")
        t0 = job["events"][0]["at"] if job["events"] else time.time()
        events = [{**e, "elapsed": round(e["at"] - t0, 1)} for e in job["events"]]
        out = {"status": job["status"], "events": events, "result": job["result"], "error": job["error"]}
        if job.get("kind") == "batch":
            out["done"] = len(job.get("done_rows", []))
            out["done_rows"] = job.get("done_rows", [])
        return out

# serve React build
DIST = ROOT / "webapp" / "frontend" / "dist"
if DIST.exists():
    app.mount("/assets", StaticFiles(directory=DIST / "assets"), name="assets")

    @app.get("/{path:path}")
    def spa(path: str):
        target = DIST / path
        if path.startswith("api/") or (path and target.exists() and target.is_file()):
            return FileResponse(target) if target.is_file() else {"detail": "not found"}
        resp = FileResponse(DIST / "index.html")
        resp.headers["Cache-Control"] = "no-store, max-age=0"
        return resp
